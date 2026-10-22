from pathlib import Path
from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"D:\project10")
SHOTS = ROOT / "artifacts" / "user-guide" / "screenshots"
OUT = ROOT / "artifacts" / "role-guides"
CROPS = OUT / "crops"
OUT.mkdir(parents=True, exist_ok=True)
CROPS.mkdir(parents=True, exist_ok=True)

NAVY = "142B36"; ORANGE = "C46C20"; PALE = "F5F8F9"; LINE = "D9D9D9"
TEXT = RGBColor(20, 43, 54); MUTED = RGBColor(91, 109, 126); ACCENT = RGBColor(196, 108, 32)

def crop(name, height):
    src = SHOTS / name; dst = CROPS / f"{Path(name).stem}-{height}.png"
    if not dst.exists():
        with Image.open(src) as im:
            im.crop((0, 0, im.width, min(height, im.height))).save(dst)
    return dst

def font(run, size=10.5, bold=False, color=TEXT, name="Microsoft YaHei"):
    run.font.name = name; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = color

def border(cell):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None: borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        n = OxmlElement(f"w:{edge}"); n.set(qn("w:val"), "single"); n.set(qn("w:sz"), "6"); n.set(qn("w:color"), LINE); borders.append(n)

def shade(cell, fill):
    n = OxmlElement("w:shd"); n.set(qn("w:fill"), fill); cell._tc.get_or_add_tcPr().append(n)

def margins(cell):
    tcMar = OxmlElement("w:tcMar")
    for side, value in (("top",120),("start",140),("bottom",120),("end",140)):
        n=OxmlElement(f"w:{side}"); n.set(qn("w:w"),str(value)); n.set(qn("w:type"),"dxa"); tcMar.append(n)
    cell._tc.get_or_add_tcPr().append(tcMar)

def page_num(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(paragraph.add_run("第 "), 8.5, False, MUTED)
    run = paragraph.add_run(); begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin")
    text=OxmlElement("w:instrText"); text.set(qn("xml:space"),"preserve"); text.text="PAGE"
    end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end")
    run._r.extend([begin,text,end]); font(paragraph.add_run(" 页"), 8.5, False, MUTED)

def setup(title, audience):
    d=Document(); sec=d.sections[0]
    sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.7); sec.left_margin=Cm(2); sec.right_margin=Cm(2)
    normal=d.styles["Normal"]; normal.font.name="Microsoft YaHei"; normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); normal.font.size=Pt(10.5); normal.font.color.rgb=TEXT
    for name,size,before,after in (("Title",27,0,16),("Heading 1",20,13,8),("Heading 2",14,9,5)):
        s=d.styles[name]; s.font.name="Microsoft YaHei"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
        pPr=s._element.get_or_add_pPr(); pBdr=pPr.find(qn("w:pBdr"));
        if pBdr is not None: pPr.remove(pBdr)
    footer=sec.footer.paragraphs[0]; font(footer.add_run(f"天合教育 IT 运维平台  {audience}使用手册    "),8.5,False,MUTED); page_num(footer)
    logo=ROOT/"public"/"brand"/"tianhe-education-ops-symbol-v2.png"
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(70); p.add_run().add_picture(str(logo),width=Cm(3.1))
    p=d.add_paragraph(title,style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(f"适用对象  {audience}"),12.5,True,ACCENT)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(12); font(p.add_run("企业微信工作台内使用  版本日期 2026 年 9 月"),9.5,False,MUTED)
    return d

def h(d,text,level=1):
    p=d.add_paragraph(text,style=f"Heading {level}"); p.paragraph_format.keep_with_next=True; return p

def body(d,text,lead=None):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(7); p.paragraph_format.line_spacing=1.35
    if lead and text.startswith(lead): font(p.add_run(lead),10.5,True,TEXT); font(p.add_run(text[len(lead):]),10.5,False,TEXT)
    else: font(p.add_run(text),10.5,False,TEXT)
    return p

def steps(d,items):
    for i,item in enumerate(items,1):
        p=d.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.25
        font(p.add_run(f"{i}. "),10.5,True,ACCENT); font(p.add_run(item),10.5,False,TEXT)

def bullets(d,items):
    for item in items:
        p=d.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.25; font(p.add_run(item),10.5,False,TEXT)

def fig(d,path,caption,width=16):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True; p.add_run().add_picture(str(path),width=Cm(width))
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(9); font(p.add_run(caption),8.5,False,MUTED)

def page(d): d.add_page_break()

def status_table(d):
    rows=[("待接单","已进入队列，等待运维接单"),("已接单或处理中","运维正在处理"),("等待用户","请员工补充信息或测试"),("等待配件","正在等待设备或外部条件"),("待员工确认","运维已处理，请员工确认结果"),("已完成","员工确认问题已解决"),("已撤回","员工主动撤回，不计完成量"),("已取消","运维取消，不计完成量")]
    t=d.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,txt,w in zip(t.rows[0].cells,["状态","含义"],[Cm(4.2),Cm(12.2)]):
        c.width=w; shade(c,NAVY); border(c); margins(c); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(txt),9.5,True,RGBColor(255,255,255))
    for n,(state,meaning) in enumerate(rows):
        cells=t.add_row().cells
        for i,(c,txt,w) in enumerate(zip(cells,[state,meaning],[Cm(4.2),Cm(12.2)])):
            c.width=w; border(c); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if n%2: shade(c,PALE)
            p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT; font(p.add_run(txt),9.4,i==0,TEXT)

# Employee manual
d=setup("天合教育 IT 运维平台普通员工使用手册","普通员工")
body(d,"本手册帮助员工从企业微信进入 IT 服务台，完成报修、查看进度、与运维沟通、撤回误报工单，并在问题解决后确认结果。所有说明文字均可在 Word 中直接修改。")
page(d); h(d,"1 进入系统和查看工单")
h(d,"1.1 从企业微信进入",2); steps(d,["打开企业微信并确认当前企业为“天合教育”。","进入工作台，点击“天合运维平台”或“IT 服务台”。","系统自动识别姓名和部门，不需要注册密码。"])
h(d,"1.2 我的工单",2); body(d,"首页“我的工单”只显示自己提交的内容。点击工单可以查看排队、处理时间线、聊天记录和维修结果，其他员工不能通过修改网址查看您的工单。")
fig(d,crop("01-employee-my-tickets.png",844),"图 1  我的工单列表",6.4)
page(d); h(d,"2 提交报修")
steps(d,["点击底部“提交报修”。","核对部门和联系电话。填写故障设备、故障地点、问题分类、问题标题和问题描述。","需要时上传截图。最多 5 张，支持 JPG、PNG、WebP、GIF，单张不超过 5 MB。","一般问题选“普通”；只有业务中断等紧急情况选“紧急”，并填写紧急原因。","点击“提交报修工单”。如果信息不完整，系统会提示且保留已经填写的内容。"])
body(d,"描述示例：今天 09:20，2616 教室的 Mac Mini 无法开机，已经重新插拔电源并更换插座，仍无反应。")
fig(d,crop("02-employee-new-ticket.png",844),"图 2  手机端报修表单",6.3)
page(d); h(d,"3 查看进度和与运维沟通")
steps(d,["在“我的工单”点击需要查看的工单。","查看当前排队位置和维修时间线。","在“工单沟通”中补充文字或图片。运维回复后会在列表显示新消息。","状态为“等待用户”时，请尽快回复运维需要的信息。","运维处理后，工单进入“待员工确认”。确认已解决后点击“确认已解决”；如果仍有问题，选择“问题未解决”并说明情况。"])
fig(d,crop("03-employee-ticket-detail.png",844),"图 3  工单详情和处理进度",6.3)
page(d); h(d,"4 撤回工单和意见建议")
h(d,"4.1 撤回工单",2); body(d,"误报、重复提交或问题自行恢复时，可以在允许撤回的状态下点击“撤回工单”，并填写撤回原因。撤回记录会保留，但不计入运维完成量。")
h(d,"4.2 解决方案",2); body(d,"提交报修前可以进入“解决方案”，输入故障现象或关键词，查看已经验证过的解决方法。没有匹配结果时再提交工单。")
fig(d,SHOTS/"08-knowledge.png","图 4  常见问题与解决方案",15.7)
h(d,"4.3 意见建议",2); body(d,"“意见建议”用于反馈服务流程、程序体验或功能需求，不会进入运维工单队列。建议内容只有管理员能看到，运维人员看不到。")
page(d); h(d,"5 工单状态说明")
status_table(d)
h(d,"6 常见问题")
for q,a in [("企业微信里打不开怎么办？","确认当前登录的是天合教育企业，彻底关闭应用页面后重新从工作台进入。"),("提交后没有消息怎么办？","先在“我的工单”查看是否已经生成工单编号；如果已生成，说明提交成功。"),("图片不显示怎么办？","刷新工单详情；仍不显示时记录工单号并通过聊天告知运维。")]:
    body(d,q); body(d,a)
d.save(OUT/"天合教育IT运维平台普通员工使用手册.docx")

# Operations manual
d=setup("天合教育 IT 运维平台运维人员使用手册","IT 运维人员")
body(d,"本手册用于运维人员日常接单、沟通、记录、状态流转和知识沉淀。运维账号同时保留员工能力，因此也可以提交自己的报修工单。所有说明文字均可在 Word 中直接修改。")
page(d); h(d,"1 今日工作台")
body(d,"工作台顶部显示待处理、处理中、今日完成和 SLA 超时数量。下方队列按优先级和排序展示所有未完成工单；紧急工单提高优先级，但原始提交时间不会改变。")
steps(d,["先查看带“紧急”或“SLA 已超时”标识的工单。","点击工单，核对员工、部门、设备、地点、故障描述和截图。","只有确实开始负责时才接单；无法处理时应转交给合适的运维人员。"])
fig(d,crop("05-ops-workspace.png",900),"图 1  今日工作台和工单队列",15.8)
page(d); h(d,"2 处理一张工单")
steps(d,["点击“接单”，系统记录接单人和接单时间，并通知员工。","点击“开始处理”，根据进度添加检查记录和处理截图。","通过“工单沟通”询问员工或发送文字解决方案。员工回复后，工作台显示新消息。","需要员工配合时切换为“等待用户”；等待设备时切换为“等待配件”。","处理结束必须填写故障原因、处理方式和处理结果。完成后状态变为“待员工确认”。","员工确认解决后才变成“已完成”；员工反馈未解决时，工单重新进入处理队列。"])
fig(d,crop("06-ops-ticket-detail.png",900),"图 2  工单详情、沟通和操作区",15.8)
page(d); h(d,"3 指派 转交 取消和优先级")
bullets(d,["指派：为未明确负责人的工单指定运维人员。","转交：处理中更换负责人，必须填写原因，时间线保留完整记录。","取消：仅用于不应继续处理的工单，必须填写取消原因。运维取消与员工撤回分别统计。","优先级：根据实际影响调整普通或紧急，不修改原始提交时间。"])
h(d,"4 历史工单")
body(d,"历史工单支持按工单号、姓名、部门、设备和关键词搜索，也可以组合状态、分类、优先级和日期范围。常用筛选条件可以命名保存。")
fig(d,crop("07-ops-history.png",1000),"图 3  历史工单搜索和组合筛选",15.8)
page(d); h(d,"5 处理记录标准")
body(d,"处理记录应让另一位运维人员无需口头说明也能继续处理。建议按照“检查了什么、发现了什么、做了什么、结果如何”的顺序填写。")
bullets(d,["不合格：已处理、好了、重启。","合格：检查交换机 3 号端口，发现端口无链路；更换至 5 号端口后连通，连续测试 10 分钟正常。","聊天用于即时沟通；关键技术动作仍要写入处理记录。","不得记录账号密码、企业微信 Secret、服务器私钥等敏感信息。"])
h(d,"6 发布知识记录")
body(d,"只有填写完整维修结论的工单才能加入知识库。发布前确认故障现象、原因和解决方法清晰，并补充检索关键词。知识内容面向员工自助查询，不得包含敏感信息。")
fig(d,SHOTS/"08-knowledge.png","图 4  面向员工的解决方案页面",15.8)
page(d); h(d,"7 状态使用原则")
status_table(d)
h(d,"8 企业微信提醒")
body(d,"新工单、员工回复等提醒是否发送给某位运维，由管理员在“成员权限”中配置。收到提醒后应进入系统查看完整内容，不要只依据通知摘要处理。")
d.save(OUT/"天合教育IT运维平台运维人员使用手册.docx")

# Administrator manual
d=setup("天合教育 IT 运维平台管理员使用手册","管理员和领导")
body(d,"本手册用于查看真实工单统计、管理成员角色和提醒、维护 SLA 与值班安排、查看员工建议和系统运行状态。管理员默认只查看和配置，不参与接单；需要处理工单时必须另外授予运维角色。所有说明文字均可在 Word 中直接修改。")
page(d); h(d,"1 数据看板")
body(d,"数据看板来源于真实工单和状态时间戳。可以选择今天、本周、本月或自定义日期，并导出 Excel。待员工确认的工单尚未计入完成量；员工撤回和运维取消分别展示，均不计入完成量。")
fig(d,crop("09-admin-dashboard.png",1000),"图 1  IT 运维数据看板",15.8)
bullets(d,["期间报修：所选时间内新建的工单数。","期间完成：员工在所选时间内确认已解决的工单数。","平均响应时间：提交至首次接单。","平均解决时间：开始处理至员工确认完成。","取消明细：区分员工撤回和运维取消。"])
page(d); h(d,"2 成员权限")
body(d,"所有通过企业微信进入系统的成员默认拥有员工角色，可以提交自己的工单。管理员可以额外授予运维或管理员角色。")
steps(d,["进入“成员权限”，按姓名、部门或企业微信 UserID 搜索。","勾选或取消“运维”“管理员”。员工角色始终保留。","只有运维角色可以接单和回复工单；管理员角色本身不处理工单。","点击“保存权限”并确认。配置文件锁定的核心账号不能从页面移除对应权限。"])
fig(d,crop("10-admin-users.png",900),"图 2  成员角色和提醒设置",15.8)
page(d); h(d,"3 运维提醒")
body(d,"运维角色下方可以设置新工单、新回复、SLA、日报等提醒。普通员工不需要配置运维提醒，员工会自动收到自己工单的接单、状态变化、回复和待确认通知。")
bullets(d,["新工单：有人提交工单时提醒指定运维。","新回复：员工在工单聊天中回复时提醒负责运维。","SLA：工单接近或超过时限时提醒。","暂停全部：临时停止该运维的全部提醒，不删除其运维权限。"])
h(d,"4 员工建议")
body(d,"建议管理仅管理员可见。可以按待查看、已查看和已采纳筛选。运维人员不能查看建议内容。")
fig(d,SHOTS/"11-admin-suggestions.png","图 3  员工建议管理",15.8)
page(d); h(d,"5 SLA 值班和接单状态")
body(d,"“运维设置”用于维护普通和紧急工单的响应时限、解决时限、工作时段、工作日、运维可接单状态和值班安排。修改会影响队列超时判断和通知策略，应由明确负责人维护。")
fig(d,SHOTS/"12-admin-operations.png","图 4  运维规则和排班",15.8)
page(d)
h(d,"6 系统状态")
body(d,"系统状态页用于检查服务、消息队列、磁盘空间、最近备份、工单总数、附件用量和审计记录。出现磁盘告警、消息失败持续增加或长期无备份时，应联系服务器维护人员。")
fig(d,SHOTS/"13-admin-system.png","图 5  系统运行状态",15.8)
page(d); h(d,"7 管理边界和安全要求")
bullets(d,["管理员查看数据和维护配置，不应代替运维处理工单。","授予管理员权限前确认业务需要，避免无关人员看到跨部门统计。","不要在工单或建议中记录密码、Secret、私钥等敏感信息。","成员离职或职责变化时及时取消额外角色和提醒。","导出的 Excel 含内部运维数据，应按公司内部资料管理。"])
h(d,"8 常见问题")
for q,a in [("为什么管理员看不到接单按钮？","管理员默认不处理工单。如确需接单，请同时授予该账号运维角色。"),("新员工能直接使用吗？","可以。成员在自建应用可见范围内，首次进入时系统自动创建员工账号。"),("为什么某位运维收不到新单提醒？","检查其企业微信 UserID 是否已绑定、是否有运维角色、新工单提醒是否开启，以及系统状态中的消息失败数。"),("完成量为什么少一张？","运维处理完成后仍需员工确认。处于“待员工确认”的工单不计入完成量。")]:
    body(d,q); body(d,a)
d.save(OUT/"天合教育IT运维平台管理员使用手册.docx")

for p in sorted(OUT.glob("*.docx")): print(p)
