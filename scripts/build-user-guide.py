from pathlib import Path
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(r"D:\project10")
SRC = ROOT / "artifacts" / "user-guide" / "screenshots"
OUT = ROOT / "artifacts" / "user-guide" / "天合教育IT运维平台使用说明.docx"
CROPS = ROOT / "artifacts" / "user-guide" / "crops"
CROPS.mkdir(parents=True, exist_ok=True)

NAVY = "142B36"
ORANGE = "C46C20"
LIGHT_BLUE = "EAF2F5"
PALE = "F7F9FA"
LINE = "D9D9D9"
TEXT = RGBColor(20, 43, 54)
MUTED = RGBColor(91, 109, 126)

def crop_top(name, height):
    src = SRC / name
    dst = CROPS / name
    with Image.open(src) as im:
        im.crop((0, 0, im.width, min(im.height, height))).save(dst)
    return dst

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_border(cell, color=LINE):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), color)
        borders.append(tag)

def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText"); instrText.set(qn("xml:space"), "preserve"); instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)
    paragraph.add_run(" 页")

def set_font(run, size=None, bold=None, color=None, name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = color

def add_body(text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.35
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead); set_font(r1, 10.5, True, TEXT)
        r2 = p.add_run(text[len(bold_lead):]); set_font(r2, 10.5, False, TEXT)
    else:
        r = p.add_run(text); set_font(r, 10.5, False, TEXT)
    return p

def add_bullets(items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.25
        set_font(p.add_run(item), 10.5, False, TEXT)

def add_steps(items):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.15)
        p.paragraph_format.first_line_indent = Cm(-0.15)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(f"{i}. "); set_font(r, 10.5, True, RGBColor.from_string(ORANGE))
        set_font(p.add_run(item), 10.5, False, TEXT)

def add_heading(text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p

def add_figure(image_path, caption, width_cm=16.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(10)
    set_font(cp.add_run(caption), 8.5, False, MUTED)

def page_break():
    doc.add_page_break()

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.9); sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5); normal.font.color.rgb = TEXT
for style_name, size, before, after in (("Title", 27, 0, 18), ("Heading 1", 20, 14, 9), ("Heading 2", 14, 10, 6)):
    st = styles[style_name]
    st.font.name = "Microsoft YaHei"; st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
    pPr = st._element.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is not None:
        pPr.remove(pBdr)

footer = sec.footer.paragraphs[0]
set_font(footer.add_run("天合教育 IT 运维平台使用说明    "), 8.5, False, MUTED)
add_page_number(footer)

# Cover
logo = ROOT / "public" / "brand" / "tianhe-education-ops-symbol-v2.png"
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(80); p.add_run().add_picture(str(logo), width=Cm(3.2))
p = doc.add_paragraph("天合教育 IT 运维平台使用说明", style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("普通员工  运维人员  管理员"), 13, True, RGBColor.from_string(ORANGE))
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(16)
set_font(p.add_run("适用于企业微信工作台内的 IT 服务台\n版本日期 2026 年 9 月"), 10.5, False, MUTED)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(60)
set_font(p.add_run("阅读建议"), 12, True, TEXT)
add_body("第一次使用时，先阅读第 1 章了解自己的角色，再直接跳到对应章节。截图中的姓名、工单和统计数据仅用于操作演示，您在正式系统中看到的是本企业的真实数据。")

page_break()
add_heading("1 快速开始", 1)
add_heading("1.1 如何进入系统", 2)
add_steps([
    "打开企业微信，进入工作台。",
    "找到“天合运维平台”或“IT 服务台”并点击。",
    "系统会自动识别企业微信身份、姓名和部门，不需要另外注册密码。",
    "首次进入后，底部或左侧菜单会根据您的权限自动显示。",
])
add_body("如果页面偶尔无法打开：先确认企业微信当前登录的是“天合教育”企业，再彻底关闭该应用页面并重新从工作台进入。不要把企业微信内的链接复制到普通浏览器中登录。", "如果页面偶尔无法打开：")

add_heading("1.2 三种角色能做什么", 2)
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(3.0), Cm(4.5), Cm(4.5), Cm(4.8)]
headers = ["角色", "主要入口", "可以做什么", "不能做什么"]
for i, (cell, text, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
    cell.width = width; set_cell_shading(cell, NAVY); set_cell_border(cell); set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), 9.5, True, RGBColor(255,255,255))
set_repeat_table_header(table.rows[0])
rows = [
    ("普通员工", "我的工单\n提交报修\n解决方案\n意见建议", "提交、撤回自己的工单；查看进度；与运维沟通；确认是否解决", "不能查看他人工单；不能处理工单；不能查看管理数据"),
    ("运维人员", "今日工作台\n历史工单", "接单、处理、沟通、转交、完成工单；检索历史；发布知识记录", "不能查看员工建议；不能修改成员管理员权限"),
    ("管理员", "数据看板\n建议管理\n成员权限\n运维设置\n系统状态", "查看统计和取消明细；管理员工权限、提醒、值班和系统状态", "默认不接单、不处理工单；只有同时授予运维角色才可处理"),
]
for ri, row in enumerate(rows):
    cells = table.add_row().cells
    for i, (cell, text, width) in enumerate(zip(cells, row, widths)):
        cell.width = width; set_cell_border(cell); set_cell_margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if ri % 2: set_cell_shading(cell, PALE)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(text), 9.2, i == 0, TEXT)

add_heading("1.3 一个账号有多个角色时", 2)
add_body("一个企业微信账号可以同时拥有员工、运维和管理员权限，不需要切换企业微信账号，也没有单独的“切换角色”按钮。系统会把所有已授权菜单同时显示出来。例如，张勇可以先从“提交报修”创建自己的工单，再从“今日工作台”处理工单，还可以从“数据看板”查看统计。")

add_heading("2 普通员工使用说明", 1)
add_heading("2.1 查看我的工单", 2)
add_steps([
    "进入系统后点击底部“我的工单”。",
    "查看工单编号、标题、优先级和当前状态。",
    "点击任意一张工单进入详情页，查看排队、时间线、运维回复和处理结果。",
])
add_figure(crop_top("01-employee-my-tickets.png", 844), "图 1  普通员工的“我的工单”页面", 7.5)

page_break()
add_heading("2.2 提交报修", 2)
add_steps([
    "点击“提交报修”。部门和联系电话会尽量自动带出，请核对是否正确。",
    "填写设备、地点、问题分类、标题和详细描述。描述应包含出现时间、故障现象和已经尝试过的操作。",
    "需要时上传故障截图。图片支持 JPG、PNG、WebP、GIF，单张不超过 5 MB。",
    "一般问题选择“普通”。只有业务中断等情况选择“紧急”，并填写紧急原因。",
    "点击“提交报修工单”。如有必填项遗漏，页面会提示，已经填写的内容会保留。",
])
add_figure(crop_top("02-employee-new-ticket.png", 844), "图 2  手机端报修表单", 7.5)

page_break()
add_heading("2.3 查看进度和与运维沟通", 2)
add_body("工单详情页会显示报修信息、维修时间线和沟通记录。运维接单、状态变化、回复和处理完成时，系统会按配置发送企业微信应用消息。")
add_steps([
    "打开工单详情，在“工单沟通”中输入补充信息或回复，可附带图片。",
    "如果运维把工单设为“等待用户”，请尽快回复所需信息。",
    "运维处理完后，工单会进入“待员工确认”。确认已经解决后点击“确认已解决”，此时工单才计为“已完成”。",
    "如果问题仍存在，选择“问题未解决”并说明情况，工单会重新回到处理队列。",
])
add_figure(crop_top("03-employee-ticket-detail.png", 844), "图 3  员工工单详情和处理进度", 6.0)

page_break()
add_heading("2.4 撤回工单", 2)
add_body("如果误报、重复提交或问题自行恢复，可在允许撤回的状态下使用“撤回工单”。撤回时需要填写原因。撤回后运维和管理员仍能看到撤回记录，但不会计入完成量。")
add_heading("2.5 查找解决方案", 2)
add_body("点击“解决方案”，输入故障现象或关键词。知识库内容来源于已完成工单，适合先自助排查常见问题。如果没有匹配结果，再提交报修。")
add_figure(SRC / "08-knowledge.png", "图 4  常见问题与解决方案页面", 16.0)
add_heading("2.6 提交意见建议", 2)
add_body("“意见建议”用于反馈服务流程、程序体验或功能需求，不会进入运维工单队列。建议内容只有管理员能看到，运维人员无权查看。")

page_break()
add_heading("3 运维人员使用说明", 1)
add_heading("3.1 今日工作台", 2)
add_body("进入“今日工作台”后，先看顶部的待处理、处理中、今日完成和 SLA 超时数量，再按队列顺序处理。紧急工单会提高排序优先级，但原始提交时间不会改变。")
add_steps([
    "优先查看带“紧急”或“SLA 已超时”标识的工单。",
    "点击工单进入详情，先核对报修人、部门、设备、地点和描述。",
    "接单后及时开始处理；需要用户补充信息时选择“等待用户”；缺少配件时选择“等待配件”。",
])
add_figure(crop_top("05-ops-workspace.png", 900), "图 5  运维今日工作台和当前队列", 16.0)

page_break()
add_heading("3.2 处理一张工单", 2)
add_steps([
    "点击“接单”，系统记录接单人和时间，并通知员工。",
    "点击“开始处理”，在处理过程中随时添加检查记录和处理截图。",
    "使用“工单沟通”向员工询问情况或发送文字解决方法；员工回复后工作台会显示新消息。",
    "完成处理时填写故障原因、处理方式和处理结果。运维确认处理完成后，状态变为“待员工确认”。",
    "员工确认解决后，工单才进入“已完成”并纳入完成量；若员工反馈未解决，工单重新进入处理流程。",
])
add_figure(crop_top("06-ops-ticket-detail.png", 900), "图 6  运维工单详情、沟通和工单操作", 16.0)

page_break()
add_heading("3.3 历史工单检索", 2)
add_body("“历史工单”支持按工单号、姓名、部门、设备和关键词搜索，也可按状态、分类、优先级、日期范围和排序方式筛选。常用条件可以命名保存。")
add_figure(crop_top("07-ops-history.png", 1000), "图 7  历史工单搜索和筛选", 16.0)
add_heading("3.4 指派 转交和取消", 2)
add_bullets([
    "指派：把尚未明确负责人的工单交给某位运维人员。",
    "转交：处理过程中更换负责人员，必须填写原因，时间线会保留记录。",
    "取消：只用于工单不应继续处理的情况，必须填写取消原因。运维取消与员工撤回会分别统计。",
])
add_heading("3.5 发布知识记录", 2)
add_body("填写完整维修结论的工单可以发布到知识库。发布前检查故障现象、原因和解决方法是否清楚，避免写入手机号、账号密码或其他敏感信息。")

page_break()
add_heading("4 管理员使用说明", 1)
add_heading("4.1 查看数据看板", 2)
add_body("管理员从“数据看板”查看真实工单统计。可选择今天、本周、本月或自定义日期，并导出 Excel。待员工确认的工单尚未计入完成量；员工撤回和运维取消分别展示，不计入完成量。")
add_figure(crop_top("09-admin-dashboard.png", 1000), "图 8  管理员数据看板", 16.0)

page_break()
add_heading("4.2 管理成员权限", 2)
add_body("所有通过企业微信进入系统的成员默认都是员工，可正常提交报修。管理员可以额外授予“运维”或“管理员”权限。保存后权限立即生效，并按系统配置发送提醒。")
add_steps([
    "进入“成员权限”，搜索姓名、部门或企业微信 UserID。",
    "勾选需要增加的“运维”或“管理员”角色。普通员工身份始终保留。",
    "点击“保存权限”并确认。配置文件锁定的核心账号不能在页面中取消对应权限。",
    "管理员本身不处理工单；确需处理时，必须同时授予运维角色。",
])
add_figure(crop_top("10-admin-users.png", 900), "图 9  成员角色和提醒权限管理", 16.0)

page_break()
add_heading("4.3 查看员工建议", 2)
add_body("进入“建议管理”查看员工提交的服务建议、程序建议和功能需求。管理员可以标记为已查看或已采纳；运维角色看不到此页面和建议内容。")
add_figure(SRC / "11-admin-suggestions.png", "图 10  仅管理员可见的员工建议", 16.0)
page_break()
add_heading("4.4 运维设置", 2)
add_body("“运维设置”用于维护 SLA 响应和解决时限、工作时段、运维可接单状态和值班安排。修改后会影响超时判断和队列提醒，应由明确负责人维护。")
add_figure(SRC / "12-admin-operations.png", "图 11  SLA 与值班设置", 16.0)

page_break()
add_heading("4.5 系统状态", 2)
add_body("“系统状态”是只读检查页，用于查看服务是否正常、消息队列失败数、磁盘空间、最近备份、工单总数和附件用量。发现“磁盘空间不可读”“暂无备份”或消息失败持续增加时，应联系服务器维护人员处理。")
add_figure(SRC / "13-admin-system.png", "图 12  系统运行状态", 16.0)

page_break()
add_heading("5 工单状态怎么理解", 1)
status_table = doc.add_table(rows=1, cols=3)
status_table.alignment = WD_TABLE_ALIGNMENT.CENTER
status_table.autofit = False
sw = [Cm(3.2), Cm(7.2), Cm(6.4)]
for cell, text, width in zip(status_table.rows[0].cells, ["状态", "表示什么", "下一步由谁做"], sw):
    cell.width = width; set_cell_shading(cell, NAVY); set_cell_border(cell); set_cell_margins(cell)
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run(text), 9.5, True, RGBColor(255,255,255))
set_repeat_table_header(status_table.rows[0])
statuses = [
    ("待接单", "工单已经进入队列，还没有运维接单", "运维接单"),
    ("已接单", "运维已负责，但尚未正式开始处理", "运维开始处理"),
    ("处理中", "运维正在检查或实施解决方案", "运维更新记录"),
    ("等待用户", "需要员工补充信息、测试或回复", "员工回复"),
    ("等待配件", "需要等待设备、耗材或外部条件", "运维跟进"),
    ("待员工确认", "运维已处理，等待员工确认结果", "员工确认已解决或反馈未解决"),
    ("已完成", "员工已经确认问题解决", "流程结束并计入完成统计"),
    ("已取消", "由运维取消，不再继续处理", "保留取消原因，不计入完成量"),
    ("已撤回", "员工主动撤回自己的工单", "保留撤回原因，不计入完成量"),
]
for ri, row in enumerate(statuses):
    cells = status_table.add_row().cells
    for i, (cell, text, width) in enumerate(zip(cells, row, sw)):
        cell.width = width; set_cell_border(cell); set_cell_margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if ri % 2: set_cell_shading(cell, PALE)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(text), 9.3, i == 0, TEXT)

add_heading("6 常见问题", 1)
faqs = [
    ("新员工第一次进入能直接报修吗？", "可以。只要成员在企业微信自建应用的可见范围内，第一次进入时系统会自动创建员工账号，默认拥有提交报修权限。"),
    ("为什么我看不到今日工作台？", "该账号尚未被授予运维角色。请联系管理员在“成员权限”中添加运维权限。"),
    ("为什么管理员看不到接单按钮？", "管理员默认只查看数据，不参与处理。需要接单时，还必须同时拥有运维角色。"),
    ("企业微信消息没收到怎么办？", "先确认自建应用对该成员可见、成员 UserID 已绑定、应用消息未被关闭，再由管理员查看系统状态中的消息失败数。"),
    ("上传图片后看不到怎么办？", "刷新工单详情重新查看。如果仍不显示，记录工单号并联系运维检查附件存储和访问权限。"),
    ("提交时提示信息不完整怎么办？", "按照红色提示补充必填项。系统会保留已填写内容，不需要从头重新填写。"),
]
for q, a in faqs:
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(7); p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(q), 10.5, True, TEXT)
    add_body(a)

add_heading("7 使用时请记住", 1)
add_bullets([
    "不要在工单、聊天或知识库中填写账号密码、企业微信 Secret、服务器私钥等敏感信息。",
    "紧急优先级只用于业务中断或大范围影响，并要写清影响范围。",
    "处理过程要持续记录，完成时填写清楚故障原因、处理方式和结果。",
    "员工确认已解决后工单才真正完成；如果未解决，应立即反馈并重新进入处理流程。",
    "遇到系统页面无法打开、消息长期未送达或磁盘告警时，联系管理员或服务器维护人员。",
])

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
