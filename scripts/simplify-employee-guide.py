from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

src = Path(r"C:\Users\Administrator\Desktop\天合教育IT运维平台普通员工使用手册.docx")
out = Path(r"D:\project10\artifacts\role-guides\天合教育IT运维平台普通员工使用手册-大白话版.docx")

doc = Document(src)

replacements = {
    2: "适用对象  公司普通员工",
    3: "在企业微信工作台里使用  2026 年 9 月版",
    4: "这本手册教你怎么报修、怎么看处理进度、怎么和运维沟通。第一次用也不用担心，跟着下面的步骤操作就可以。",
    6: "1 怎么进入系统和查看工单",
    7: "1.1 从企业微信进入",
    8: "1. 打开企业微信，先确认自己进入的是“天合教育”这个企业。",
    9: "2. 点开“工作台”，找到“天合运维平台”或“IT 服务台”。",
    10: "3. 点进去就能直接使用，不需要另外注册，也不用输入密码。",
    11: "1.2 查看我报过的工单",
    12: "进入首页后，你会看到“我的工单”。这里都是你自己提交的报修，别人看不到你的具体内容。点开任意一张工单，就能查看现在排到第几位、运维处理到哪一步，以及最后是怎么解决的。",
    16: "2 怎么提交报修",
    17: "1. 点击底部的“提交报修”。",
    18: "2. 按照页面提示填写。设备可以写“Mac Mini”“联想笔记本”，地点可以写“2616 教室”。问题描述尽量说清楚什么时候出现、有什么现象。",
    19: "3. 有故障截图就上传，最多 5 张。支持 JPG、PNG、WebP、GIF，每张不超过 5 MB。没有截图也可以提交。",
    20: "4. 一般问题请选择“普通”。只有确实会影响上课、直播或大面积办公时才选“紧急”，并简单写明为什么着急。",
    21: "5. 填好后点击“提交报修”。如果还有必填项没填，系统会告诉你缺少什么，前面已经填好的内容不会丢。",
    22: "可以这样写：今天 09:20，2616 教室的 Mac Mini 无法开机，已经重新插过电源，还是没有反应。",
    26: "3 怎么看进度和联系运维",
    27: "1. 在“我的工单”里点开你想查看的工单。",
    28: "2. 页面会显示目前的状态、排队位置和处理记录。",
    29: "3. 如果还需要补充情况，可以直接在“工单沟通”里发文字或图片。运维回复后，你会在企业微信收到提醒。",
    30: "4. 如果状态显示“等待用户”，说明运维正在等你补充信息，请看到后尽快回复。",
    31: "5. 运维处理好后，工单会显示“待员工确认”。如果问题确实解决了，点击“确认已解决”；如果还有问题，就选择“问题仍未解决”，并说明情况。只有你确认以后，工单才会正式显示“已完成”。",
    35: "4 报错了怎么办，还有哪些实用功能",
    36: "4.1 不小心报错了，可以撤回",
    37: "如果报错了、重复提交了，或者问题自己恢复了，可以打开工单点击“撤回工单”，再写一下原因。撤回记录不会消失，运维也能看到。",
    38: "4.2 先看看有没有现成解决办法",
    39: "有些常见问题可以自己解决。提交前可以打开“常见问题与解决方案”，搜索故障关键词。如果没有找到合适的方法，再提交工单。",
    42: "4.3 想提意见也可以",
    43: "如果你觉得服务流程或系统哪里不好用，可以打开“意见建议”告诉管理员。这个内容不会进入报修队列，只有管理员能看到，运维人员看不到。",
    45: "5 各种状态是什么意思",
    46: "6 常见问题",
    47: "企业微信里打不开怎么办？",
    48: "先确认当前登录的是“天合教育”企业，然后关闭这个页面，再从企业微信工作台重新进入。",
    49: "点了提交，但不知道有没有成功怎么办？",
    50: "打开“我的工单”看看。只要已经生成 IT 开头的工单编号，就说明提交成功了，不要重复提交。",
    51: "上传的图片看不到怎么办？",
    52: "先刷新工单详情。如果还是看不到，可以在“工单沟通”里告诉运维，请运维帮你检查。",
}

for idx, text in replacements.items():
    if idx < len(doc.paragraphs):
        doc.paragraphs[idx].text = text

table_text = [
    ("状态", "代表什么"),
    ("待接单", "工单已经提交，正在等运维接单"),
    ("已接单或处理中", "运维已经看到，正在处理"),
    ("等待用户", "运维需要你补充信息或配合测试"),
    ("等待配件", "正在等设备、配件或外部支持"),
    ("待员工确认", "运维认为已经处理好，等你确认"),
    ("已完成", "你已经确认问题解决"),
    ("已撤回", "你主动撤回了工单，不算完成量"),
    ("已取消", "运维取消了工单，不算完成量"),
]
if doc.tables:
    table = doc.tables[0]
    for row, values in zip(table.rows, table_text):
        row.cells[0].text, row.cells[1].text = values

def set_run_font(run, size=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)

for style in doc.styles:
    if hasattr(style, "font"):
        style.font.name = "Microsoft YaHei"
        try:
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            style._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
            style._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        except AttributeError:
            pass

for paragraph in doc.paragraphs:
    paragraph.paragraph_format.page_break_before = False
    for br in list(paragraph._p.xpath('.//w:br')):
        if br.get(qn('w:type')) == 'page':
            br.getparent().remove(br)
    for run in paragraph.runs:
        set_run_font(run)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run)

for section in doc.sections:
    for container in (section.header, section.footer):
        for paragraph in container.paragraphs:
            for run in paragraph.runs:
                set_run_font(run)

doc.core_properties.title = "天合教育 IT 运维平台普通员工使用手册"
doc.core_properties.subject = "普通员工报修操作说明"
doc.save(out)
print(out)
